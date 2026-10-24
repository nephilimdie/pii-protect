"""Zero-retention document service: TXT/DOCX/PDF redaction and pseudonymization.

All sample documents are synthetic and generated in memory (python-docx, reportlab).
Detection is a deterministic fake wired through the real policy/detection path
(detect_protected_entities), so the endpoint semantics match /v1/anonymize.
"""

from __future__ import annotations

import builtins
import io
import json
import logging
import os
import re
import tempfile
import zipfile

import pypdfium2 as pdfium
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from fastapi import FastAPI, HTTPException
from fastapi.responses import JSONResponse
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from PIL import Image, ImageChops, ImageFilter, ImageStat

import app.documents.docx_handler  # noqa: F401  (pre-import: no lazy import during file-watch tests)
import app.documents.pdf_handler  # noqa: F401
import app.documents.txt_handler  # noqa: F401
from app.anonymization.anonymizer import DetectionResult
from app.config import settings
from app.database import get_db
from app.detection.entities import PiiEntity
from app.documents.models import REDACTION_MARK, JoinedText
from app.identity.dependencies import require_service
from app.identity.tenant import get_tenant_id
from app.routers import documents as documents_module
from app.routers._anonymize_logic import get_anonymizer

PERSON = "Mario Rossi"
FISCAL = "RSSMRA80A01H501U"
EMAIL = "mario.rossi@example.it"
DELETED_NAME = "Luigi Verdi"
AUTHOR = "Giulia Bianchi"
SAMPLE_PII = (PERSON, FISCAL, EMAIL, DELETED_NAME, AUTHOR)
_PATTERNS = {PERSON: "PERSON", FISCAL: "FISCAL_CODE", EMAIL: "EMAIL", DELETED_NAME: "PERSON", AUTHOR: "PERSON"}
SECRET_FILENAME = "contratto_mario_rossi.docx"


# ── fakes ─────────────────────────────────────────────────────────────────────

class FakeAnonymizer:
    """Deterministic detector: exact matches of the synthetic PII strings."""

    def __init__(self, fail: bool = False) -> None:
        self.fail = fail
        self.calls = 0

    def restrict_layers(self, layers):
        pass

    def detect_only(self, text, context_id, context_type, language="it"):
        self.calls += 1
        if self.fail:
            raise RuntimeError(f"model crashed on {PERSON}")
        entities = [
            PiiEntity(m.start(), m.end(), pii_type, m.group(), 0.99, "fake")
            for value, pii_type in _PATTERNS.items() for m in re.finditer(re.escape(value), text)
        ]
        return DetectionResult(sorted(entities, key=lambda e: e.start))


def _policy(**overrides):
    policy = {
        "protect_types": None, "keep_types": set(), "surrogate_types": set(), "remove_types": set(),
        "block_types": set(), "confidence_thresholds": {}, "allowlist": {}, "detection_layers": None,
        "mode": "tag", "policy_id": "domain:hr_payroll", "policy_version": "context:1|domain:1", "policy_hash": "h",
    }
    policy.update(overrides)
    return policy


class Recorder:
    def __init__(self):
        self.usage, self.audit, self.policy_calls = [], [], []
        self.policy = _policy()


@pytest.fixture
def env(monkeypatch):
    rec = Recorder()
    anonymizer = FakeAnonymizer()

    class FakePolicyService:
        def __init__(self, db, tenant_id=None):
            pass

        async def resolve(self, **kwargs):
            rec.policy_calls.append(kwargs)
            return rec.policy

    class FakeUsage:
        def __init__(self, db):
            pass

        async def ensure_within_limits(self, api_key, chars):
            pass

        async def record(self, **kwargs):
            rec.usage.append(kwargs)

    class FakeAudit:
        def __init__(self, db, ip_anonymization=True):
            pass

        async def log(self, **kwargs):
            rec.audit.append(kwargs)

    monkeypatch.setattr(documents_module, "PolicyService", FakePolicyService)
    monkeypatch.setattr(documents_module, "UsageService", FakeUsage)
    monkeypatch.setattr(documents_module, "AuditService", FakeAudit)

    app = FastAPI()

    @app.exception_handler(HTTPException)
    async def _flat(request, exc):  # same flattening as app.main.http_exception_handler
        payload = exc.detail if isinstance(exc.detail, dict) else {"detail": exc.detail}
        return JSONResponse(payload, status_code=exc.status_code)

    app.include_router(documents_module.router, prefix="/v1")

    async def _db():
        yield None

    app.dependency_overrides[require_service] = lambda: type("Key", (), {"id": "key-1", "role": "service"})()
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[get_anonymizer] = lambda: anonymizer
    app.dependency_overrides[get_tenant_id] = lambda: None
    rec.anonymizer = anonymizer
    rec.client = TestClient(app)
    return rec


def _post(rec, op, payload, filename="documento", **fields):
    return rec.client.post(f"/v1/documents/{op}", files={"file": (filename, payload)}, data=fields)


# ── synthetic documents ──────────────────────────────────────────────────────

TXT = (f"Lettera di assunzione\nIl dipendente {PERSON} (CF {FISCAL}) riceve la busta paga.\n"
       f"Contatto: {EMAIL}\nCordiali saluti").encode()


def make_docx() -> bytes:
    doc = Document()
    props = doc.core_properties
    props.author, props.last_modified_by, props.comments = AUTHOR, AUTHOR, f"bozza per {PERSON}"
    props.title = f"Contratto {PERSON}"
    paragraph = doc.add_paragraph("Il dipendente ")
    paragraph.add_run("Ma")
    paragraph.add_run("rio Ro").bold = True  # entity split across differently formatted runs
    paragraph.add_run(f"ssi, codice fiscale {FISCAL}, lavora a Milano.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Email"
    table.cell(0, 1).text = EMAIL
    doc.sections[0].header.paragraphs[0].text = f"Riservato - {PERSON}"
    doc.sections[0].footer.paragraphs[0].text = "Pagina finale - documento interno"
    commented = doc.add_paragraph("Clausola di riservatezza.")
    doc.add_comment(commented.runs, text=f"Verificare con {PERSON}", author=AUTHOR, initials="GB")
    tracked = doc.add_paragraph("Sostituto: ")
    tracked._p.append(parse_xml(
        f'<w:del {nsdecls("w")} w:id="91" w:author="{AUTHOR}" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:delText>{DELETED_NAME}</w:delText></w:r></w:del>'))
    tracked._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="92" w:author="{AUTHOR}" w:date="2026-01-01T00:00:00Z">'
        f'<w:r><w:t>nessuno</w:t></w:r></w:ins>'))
    link = doc.add_paragraph("Scrivere a ")
    rid = doc.part.relate_to(f"mailto:{EMAIL}", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>{EMAIL}</w:t></w:r></w:hyperlink>'))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


PDF_DPI = 200
PDF_FONT = ("Helvetica", 14)
PDF_LABEL = "Dipendente: "
PDF_XY = (72, 500)  # inside the page box for every rotation


def make_pdf(pages: int = 1, rotation: int = 0) -> bytes:
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    for _ in range(pages):
        pdf.setPageRotation(rotation)
        pdf.setFont(*PDF_FONT)
        pdf.drawString(*PDF_XY, f"{PDF_LABEL}{PERSON}")
        pdf.drawString(72, 450, "Mansione: impiegato amministrativo")
        pdf.showPage()
    pdf.save()
    return out.getvalue()


def make_scanned_pdf() -> bytes:
    image = Image.new("RGB", (400, 200), "white")
    out = io.BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4)
    pdf.drawImage(ImageReader(image), 72, 500, 400, 200)
    pdf.showPage()
    pdf.save()
    return out.getvalue()


def _zip_text(payload: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        return "".join(archive.read(name).decode("utf-8", "replace") for name in archive.namelist())


def _docx_visible_text(payload: bytes) -> str:
    doc = Document(io.BytesIO(payload))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    parts += [p.text for s in doc.sections for p in s.header.paragraphs + s.footer.paragraphs]
    return "\n".join(parts)


# ── pure units ───────────────────────────────────────────────────────────────

def test_joined_text_maps_global_spans_to_segment_edits():
    joined = JoinedText(["ab", "cd"])
    assert joined.text == "ab\n\ncd"
    assert joined.split(1, 6) == [(0, 1, 2), (1, 0, 2)]
    assert joined.segment_at(4) == 1
    edits = joined.edits([(1, 6, "X")])
    assert [(e.segment, e.start, e.end, e.replacement) for e in edits] == [(0, 1, 2, "X"), (1, 0, 2, "")]


def test_txt_decodes_cp1252_and_rejects_binary():
    from app.documents.models import DocumentError
    from app.documents.txt_handler import decode_text
    assert decode_text("perché".encode("cp1252")) == "perché"
    with pytest.raises(DocumentError) as err:
        decode_text(b"\x00\x01\x02")
    assert err.value.status == 415


# ── TXT ──────────────────────────────────────────────────────────────────────

def test_txt_redact_removes_pii_and_keeps_text(env):
    r = _post(env, "redact", TXT, filename="note.txt", domain="hr_payroll")
    assert r.status_code == 200, r.text
    out = r.content.decode()
    assert not any(v in out for v in SAMPLE_PII)
    assert "Il dipendente" in out and "Cordiali saluti" in out and out.count(REDACTION_MARK) == 3
    assert r.headers["content-type"].startswith("text/plain")
    assert r.headers["x-pii-entity-count"] == "3"
    assert r.headers["cache-control"] == "no-store"
    assert env.policy_calls == [{"context_type": None, "domain": "hr_payroll"}]
    assert env.usage[0]["status"] == "ok" and env.usage[0]["entities_count"] == 3
    assert env.audit[0]["action"] == "document_redact" and "document_hash" not in env.audit[0]


def test_txt_pseudonymize_returns_zip_with_mapping(env):
    payload = TXT + f"\nFirma: {PERSON}".encode()
    r = _post(env, "pseudonymize", payload)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/zip"
    with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
        assert sorted(archive.namelist()) == ["mapping.json", "pseudonymized.txt"]
        text = archive.read("pseudonymized.txt").decode()
        mapping = json.loads(archive.read("mapping.json"))
    assert PERSON not in text and text.count("[PERSON_1]") == 2 and "[FISCAL_CODE_1]" in text
    assert mapping["document"] == "pseudonymized.txt"
    assert {"token": "[PERSON_1]", "type": "PERSON", "original": PERSON} in mapping["entities"]
    assert len(mapping["entities"]) == 3


# ── DOCX ─────────────────────────────────────────────────────────────────────

def test_docx_redact_removes_pii_everywhere_and_strips_metadata(env):
    r = _post(env, "redact", make_docx(), filename=SECRET_FILENAME)
    assert r.status_code == 200, r.text
    raw = _zip_text(r.content)
    for value in SAMPLE_PII:
        assert value not in raw, value  # body, table, header, comments, tracked changes, rels, docProps
    visible = _docx_visible_text(r.content)
    assert "Il dipendente" in visible and "lavora a Milano" in visible and "Pagina finale" in visible
    assert "Sostituto: nessuno" in visible  # insertion accepted, deletion dropped
    with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
        names = archive.namelist()
    assert not any("comments" in n for n in names) and "docProps/app.xml" not in names
    doc = Document(io.BytesIO(r.content))
    props = doc.core_properties
    assert (props.author, props.last_modified_by, props.comments, props.title) == ("", "", "", "")
    runs = doc.paragraphs[0].runs  # entity split over "Ma" + bold "rio Ro" + "ssi, ..."
    assert [run.text for run in runs[:2]] == ["Il dipendente ", REDACTION_MARK]  # mark in first touched run
    assert runs[2].bold and runs[2].text == "" and runs[3].text.startswith(", codice fiscale")
    assert doc.paragraphs[0].text == f"Il dipendente {REDACTION_MARK}, codice fiscale {REDACTION_MARK}, lavora a Milano."
    assert "w:author" not in raw and "mailto:" not in raw
    assert set(r.headers["x-pii-sanitized"].split(",")) >= {"core_properties", "comments", "tracked_changes"}


def test_docx_pseudonymize_tokens_and_mapping(env):
    r = _post(env, "pseudonymize", make_docx())
    assert r.status_code == 200, r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
        document = archive.read("pseudonymized.docx")
        mapping = json.loads(archive.read("mapping.json"))
    visible = _docx_visible_text(document)
    assert PERSON not in visible and EMAIL not in visible
    assert "Riservato - [PERSON_1]" in visible and "Il dipendente [PERSON_1], codice fiscale [FISCAL_CODE_1]" in visible
    assert {e["original"] for e in mapping["entities"]} == {PERSON, FISCAL, EMAIL}


# ── PDF ──────────────────────────────────────────────────────────────────────

def _render_output(payload: bytes):
    pdf = pdfium.PdfDocument(payload)
    try:
        chars = sum(pdf[i].get_textpage().count_chars() for i in range(len(pdf)))
        image = pdf[0].render(scale=PDF_DPI / 72).to_pil().convert("L")
        return len(pdf), chars, image
    finally:
        pdf.close()


def test_pdf_redact_rasterizes_and_boxes_entity(env):
    r = _post(env, "redact", make_pdf(pages=2))
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    pages, chars, image = _render_output(r.content)
    assert pages == 2 and chars == 0  # no text layer survives
    scale = PDF_DPI / 72
    label_w = stringWidth(PDF_LABEL, *PDF_FONT)
    name_w = stringWidth(PERSON, *PDF_FONT)
    y_px = int((A4[1] - PDF_XY[1] - PDF_FONT[1] * 0.35) * scale)
    for fraction in (0.1, 0.5, 0.9):  # across the whole name
        x_px = int((PDF_XY[0] + label_w + name_w * fraction) * scale)
        assert image.getpixel((x_px, y_px)) < 40, fraction
    label_box = image.crop((int(PDF_XY[0] * scale), y_px - 10, int((PDF_XY[0] + label_w - 8) * scale), y_px + 10))
    assert ImageStat.Stat(label_box).mean[0] > 120  # label not blacked out
    assert r.headers["x-pii-entity-count"] == "2"


def _displayed(payload: bytes):
    pdf = pdfium.PdfDocument(payload)
    try:
        page = pdf[0]
        return page.get_rotation(), page.get_size(), page.render(scale=PDF_DPI / 72).to_pil().convert("L")
    finally:
        pdf.close()


def _dark(image) -> int:
    return sum(image.point(lambda v: 255 if v < 128 else 0).histogram()[255:])


@pytest.mark.parametrize("rotation", [90, 180, 270])
def test_pdf_redact_preserves_rotation_size_and_box_position(env, rotation):
    source = make_pdf(rotation=rotation)
    r = _post(env, "redact", source)
    assert r.status_code == 200, r.text
    src_rot, src_size, before = _displayed(source)
    out_rot, out_size, after = _displayed(r.content)
    assert out_rot == src_rot == rotation
    assert out_size == pytest.approx(src_size, abs=0.01)
    common = (0, 0, min(before.width, after.width), min(before.height, after.height))
    assert abs(after.width - before.width) <= 1 and abs(after.height - before.height) <= 1
    before, after = before.crop(common), after.crop(common)
    # Only the name changed, it is now a solid box, and every other glyph is still visible.
    # Pixels that turned from paper to ink; erosion drops isolated anti-aliasing/JPEG noise.
    changed = ImageChops.multiply(before.point(lambda v: 255 if v > 128 else 0),
                                  after.point(lambda v: 255 if v < 128 else 0)).filter(ImageFilter.MinFilter(5)).getbbox()
    assert changed is not None
    assert ImageStat.Stat(after.crop(changed)).mean[0] < 40
    outside = Image.new("L", before.size, 255)
    outside.paste(0, (changed[0] - 4, changed[1] - 4, changed[2] + 4, changed[3] + 4))  # undo erosion
    keep_before = ImageChops.lighter(before, ImageChops.invert(outside))
    keep_after = ImageChops.lighter(after, ImageChops.invert(outside))
    assert _dark(keep_after) == pytest.approx(_dark(keep_before), rel=0.1) and _dark(keep_before) > 500


def test_pdf_preview_reports_pages(env):
    r = _post(env, "preview", make_pdf(pages=2))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["format"] == "pdf" and body["pages"] == 2
    assert [(e["type"], e["text"], e["page"]) for e in body["entities"]] == [("PERSON", PERSON, 1), ("PERSON", PERSON, 2)]
    assert "text_layer" in body["sanitized"]


def test_pdf_without_text_layer_is_rejected(env):
    r = _post(env, "redact", make_scanned_pdf())
    assert r.status_code == 422
    assert r.json()["error"] == "PDF_NO_TEXT_LAYER"


def test_pdf_pseudonymize_not_supported(env):
    r = _post(env, "pseudonymize", make_pdf())
    assert r.status_code == 422 and r.json()["error"] == "PSEUDONYMIZE_NOT_SUPPORTED_FOR_PDF"
    assert env.anonymizer.calls == 0


def test_pdf_page_limit(env, monkeypatch):
    monkeypatch.setattr(settings, "document_upload_max_pages", 1)
    r = _post(env, "redact", make_pdf(pages=2))
    assert r.status_code == 413 and r.json()["error"] == "DOCUMENT_PAGE_LIMIT_EXCEEDED"


# ── preview + user corrections ───────────────────────────────────────────────

def test_preview_then_redact_with_user_corrected_entities(env):
    preview = _post(env, "preview", TXT).json()
    assert preview["text"] == TXT.decode()
    entities = [e for e in preview["entities"] if e["type"] != "EMAIL"]  # user un-flags the email
    start = preview["text"].index("Cordiali")
    entities.append({"start": start, "end": start + len("Cordiali"), "type": "CUSTOM"})  # user adds a term
    calls_after_preview = env.anonymizer.calls
    r = _post(env, "redact", TXT, entities=json.dumps([{k: e[k] for k in ("start", "end", "type")} for e in entities]))
    assert r.status_code == 200, r.text
    out = r.content.decode()
    assert EMAIL in out and PERSON not in out and "Cordiali" not in out
    assert env.anonymizer.calls == calls_after_preview  # corrected list replaces detection


@pytest.mark.parametrize("entities", [
    "not json", json.dumps([{"start": 0, "end": 99999, "type": "PERSON"}]),
    json.dumps([{"start": 0, "end": 5, "type": "PERSON"}, {"start": 3, "end": 8, "type": "PERSON"}]),
    json.dumps([{"start": 0, "end": 5, "type": "lower]"}]),
])
def test_invalid_user_entities_rejected(env, entities):
    r = _post(env, "redact", TXT, entities=entities)
    assert r.status_code == 422 and r.json()["error"] == "INVALID_ENTITIES"


# ── fail-closed + limits ─────────────────────────────────────────────────────

@pytest.mark.parametrize("op", ["redact", "pseudonymize", "preview"])
def test_detector_failure_fails_closed(env, op):
    env.anonymizer.fail = True
    r = _post(env, op, make_docx())
    assert r.status_code == 503
    assert r.headers["content-type"].startswith("application/json")
    assert r.json() == {"error": "PII_PROCESSING_FAILED", "safe": False,
                        "message": "The document could not be safely processed; no document was returned."}
    assert env.usage[-1]["status"] == "error" and not env.audit


def test_blocked_type_rejects_document(env):
    env.policy = _policy(block_types={"FISCAL_CODE"})
    r = _post(env, "redact", TXT)
    assert r.status_code == 422 and r.json()["error"] == "PII_BLOCKED"
    assert FISCAL not in r.text


def test_timeout_returns_504(env, monkeypatch):
    import time as _time
    monkeypatch.setattr(settings, "document_timeout_seconds", 0.2)
    original = env.anonymizer.detect_only
    env.anonymizer.detect_only = lambda *a, **k: (_time.sleep(0.6), original(*a, **k))[1]
    r = _post(env, "redact", TXT)
    assert r.status_code == 504 and r.json()["error"] == "DOCUMENT_PROCESSING_TIMEOUT"


def test_upload_size_limit(env, monkeypatch):
    monkeypatch.setattr(settings, "document_upload_max_bytes", 2048)
    r = _post(env, "redact", b"a" * 4096)
    assert r.status_code == 413 and r.json()["error"] == "DOCUMENT_TOO_LARGE"


@pytest.mark.parametrize("payload,status,code", [
    (b"\x00\x01binary", 415, "UNSUPPORTED_DOCUMENT_TYPE"),
    (b"\xd0\xcf\x11\xe0legacy", 422, "LEGACY_OR_ENCRYPTED_OFFICE_FILE"),
    (b"", 422, "EMPTY_DOCUMENT"),
    (b"%PDF-1.7 broken", 422, "INVALID_DOCUMENT"),
])
def test_rejected_inputs(env, payload, status, code):
    r = _post(env, "redact", payload)
    assert (r.status_code, r.json()["error"]) == (status, code)


def test_requires_multipart_file(env):
    assert env.client.post("/v1/documents/redact", content=TXT).json()["error"] == "MULTIPART_REQUIRED"
    r = env.client.post("/v1/documents/redact", data={"domain": "hr_payroll"}, files={"other": ("x", b"y")})
    assert r.status_code == 422 and r.json()["error"] == "FILE_REQUIRED"


# ── zero retention ───────────────────────────────────────────────────────────

_WRITE_FLAGS = os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_APPEND | os.O_TRUNC


def test_processing_creates_no_files(env, monkeypatch, tmp_path):
    big_txt = TXT + b"\n" + b"x" * (2 * 1024 * 1024)  # above Starlette's 1 MB spool threshold
    docs = [("redact", big_txt), ("pseudonymize", make_docx()), ("redact", make_pdf()), ("preview", make_docx())]
    opened: list[str] = []
    real_os_open, real_open = os.open, builtins.open

    def watch_os_open(path, flags, *args, **kwargs):
        if flags & _WRITE_FLAGS:
            opened.append(str(path))
        return real_os_open(path, flags, *args, **kwargs)

    def watch_open(file, mode="r", *args, **kwargs):
        if isinstance(mode, str) and any(c in mode for c in "wax+"):
            opened.append(str(file))
        return real_open(file, mode, *args, **kwargs)

    monkeypatch.setattr(documents_module, "MAX_TEXT_CHARS", 3_000_000)
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    monkeypatch.setattr(os, "open", watch_os_open)
    monkeypatch.setattr(builtins, "open", watch_open)
    for op, payload in docs:
        assert _post(env, op, payload).status_code == 200
    assert opened == [] and list(tmp_path.iterdir()) == []


def test_logs_contain_no_content_or_filenames(env, caplog):
    caplog.set_level(logging.DEBUG)
    for op, payload in [("redact", make_docx()), ("pseudonymize", TXT), ("preview", make_pdf()), ("redact", make_pdf())]:
        assert _post(env, op, payload, filename=SECRET_FILENAME).status_code == 200
    env.anonymizer.fail = True
    assert _post(env, "redact", TXT, filename=SECRET_FILENAME).status_code == 503
    logged = "\n".join(f"{r.getMessage()} {r.exc_text or ''}" for r in caplog.records)
    assert "document redact ok" in logged and "failed closed: RuntimeError" in logged
    for value in (*SAMPLE_PII, SECRET_FILENAME, "contratto_mario"):
        assert value not in logged
